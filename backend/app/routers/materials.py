import os
import uuid
from datetime import datetime
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, status
from sqlalchemy.orm import Session
from app.database import get_db
from app.auth.dependencies import get_current_user
from app.config import settings
from app.models.user import User
from app.models.learning import LearningMaterial
from app.schemas.material import MaterialResponse

router = APIRouter(prefix="/materials", tags=["Learning Materials"])

ALLOWED_EXTENSIONS = {"pdf", "pptx", "ppt"}


def _extract_text(filepath: str, file_type: str) -> str:
    """Extract text from PDF or PPTX files."""
    text = ""
    try:
        if file_type == "pdf":
            from PyPDF2 import PdfReader
            reader = PdfReader(filepath)
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
        elif file_type in ("pptx", "ppt"):
            from pptx import Presentation
            prs = Presentation(filepath)
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        text += shape.text + "\n"
    except Exception as e:
        text = f"[Text extraction failed: {str(e)}]"
    return text.strip()


@router.post("/upload", response_model=MaterialResponse)
async def upload_material(
    file: UploadFile = File(...),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Upload a PDF or PPTX learning material."""
    if not file.filename:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="No file provided")

    ext = file.filename.rsplit(".", 1)[-1].lower() if "." in file.filename else ""
    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Unsupported file type. Allowed: {', '.join(ALLOWED_EXTENSIONS)}",
        )

    # Save file to disk
    os.makedirs(settings.UPLOAD_DIR, exist_ok=True)
    unique_name = f"{uuid.uuid4().hex}_{file.filename}"
    filepath = os.path.join(settings.UPLOAD_DIR, unique_name)

    content = await file.read()
    with open(filepath, "wb") as f:
        f.write(content)

    # Extract text
    extracted = _extract_text(filepath, ext)
    file_status = "processed" if extracted and not extracted.startswith("[Text extraction failed") else "failed"
    if not extracted:
        file_status = "processed"
        extracted = ""

    material = LearningMaterial(
        user_id=current_user.id,
        filename=unique_name,
        original_name=file.filename,
        file_type=ext,
        file_size=len(content),
        extracted_text=extracted,
        status=file_status,
        uploaded_at=datetime.utcnow(),
    )
    db.add(material)
    db.commit()
    db.refresh(material)

    preview = extracted[:300] + "..." if len(extracted) > 300 else extracted

    return MaterialResponse(
        id=material.id,
        filename=material.filename,
        original_name=material.original_name,
        file_type=material.file_type,
        file_size=material.file_size,
        status=material.status,
        extracted_text_preview=preview,
        uploaded_at=material.uploaded_at,
    )


@router.get("", response_model=list[MaterialResponse])
def list_materials(
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """List all materials uploaded by the current user."""
    materials = (
        db.query(LearningMaterial)
        .filter(LearningMaterial.user_id == current_user.id)
        .order_by(LearningMaterial.uploaded_at.desc())
        .all()
    )
    result = []
    for m in materials:
        preview = (m.extracted_text[:300] + "...") if m.extracted_text and len(m.extracted_text) > 300 else (m.extracted_text or "")
        result.append(
            MaterialResponse(
                id=m.id,
                filename=m.filename,
                original_name=m.original_name,
                file_type=m.file_type,
                file_size=m.file_size,
                status=m.status,
                extracted_text_preview=preview,
                uploaded_at=m.uploaded_at,
            )
        )
    return result
